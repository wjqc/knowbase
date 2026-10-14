"""extraction.parsers 委托 knowbase.parsers registry 的回归测试。

P0-6：消除解析器双实现。旧入口 parse_document() 必须走新 registry，
因此 extract / import extract 链路能拿到 PDF、DOCX、XLSX、PPTX、OCR 等能力；
同时保持旧 API 契约（ParsedDocument 结构、format、source_path、标题规则）。
"""
from __future__ import annotations

from pathlib import Path

import pytest

from knowbase.extraction.parsers import (
    ParsedDocument,
    detect_format,
    importable_extensions,
    list_supported_formats,
    parse_document,
    parse_html,
    register_parser,
)


# ---------------------------------------------------------------------------
# 分发顺序：自定义注册 > 内置 registry > 纯文本兜底
# ---------------------------------------------------------------------------


def test_custom_parser_takes_precedence_over_builtin(tmp_path: Path):
    def custom(path: Path) -> ParsedDocument:
        return ParsedDocument(text="custom body", meta={"custom": True})

    register_parser([".bridgecustom"], custom)

    f = tmp_path / "a.bridgecustom"
    f.write_text("ignored", encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta.get("custom") is True
    assert doc.text == "custom body"


def test_custom_parser_failure_returns_none(tmp_path: Path):
    def boom(path: Path) -> ParsedDocument:
        raise RuntimeError("解析失败")

    register_parser([".bridgeboom"], boom)

    f = tmp_path / "a.bridgeboom"
    f.write_text("x", encoding="utf-8")

    assert parse_document(f) is None


# ---------------------------------------------------------------------------
# 内置格式统一走 registry（meta.parser 标记实际使用的实现）
# ---------------------------------------------------------------------------


def test_markdown_delegates_to_registry(tmp_path: Path):
    f = tmp_path / "guide.md"
    f.write_text("# 标题\n\n正文内容", encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "markdown"
    assert doc.format == "markdown"
    assert doc.source_path == str(f)
    assert doc.meta["headings"] == ["标题"]
    assert len(doc.segments) >= 1


def test_markdown_title_keeps_legacy_stem_contract(tmp_path: Path):
    """旧契约：非 HTML 格式的 title 取文件名，不取文档内首个标题。"""
    f = tmp_path / "test.md"
    f.write_text("# My Document\n\nContent here", encoding="utf-8")

    doc = parse_document(f)
    assert doc.meta["title"] == "test"


def test_code_file_uses_code_parser(tmp_path: Path):
    f = tmp_path / "mod.py"
    f.write_text("def hello():\n    return 1\n", encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "code"
    assert doc.meta["language"] == "python"
    assert doc.format == "python"
    assert doc.meta["title"] == "mod"


def test_log_uses_log_parser_not_txt(tmp_path: Path):
    f = tmp_path / "app.log"
    f.write_text("2026-01-01 ERROR boom\n2026-01-01 INFO ok\n", encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "log"
    assert doc.format == "log"


def test_txt_format_stays_plaintext(tmp_path: Path):
    f = tmp_path / "note.txt"
    f.write_text("first line\nsecond", encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    # 内置 txt parser 的 meta.format 是 "plain"，对外仍用 detect_format 的标识
    assert doc.format == "plaintext"
    assert doc.meta["title"] == "note"


def test_html_drops_page_chrome_and_keeps_title(tmp_path: Path):
    f = tmp_path / "page.html"
    f.write_text(
        "<html><head><title>T</title></head>"
        "<body><nav>MENU</nav><footer>FOOT</footer>"
        "<h1>Hello</h1><p>World</p></body></html>",
        encoding="utf-8",
    )

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "html"
    assert doc.format == "html"
    assert doc.meta["title"] == "T"
    assert "MENU" not in doc.text
    assert "FOOT" not in doc.text
    assert "Hello" in doc.text


def test_html_title_falls_back_to_first_h1(tmp_path: Path):
    f = tmp_path / "no_title.html"
    f.write_text("<html><body><h1>Page Title</h1><p>x</p></body></html>", encoding="utf-8")

    doc = parse_document(f)
    assert doc.meta["title"] == "Page Title"


def test_html_title_falls_back_to_stem(tmp_path: Path):
    f = tmp_path / "plain.html"
    f.write_text("<html><body><p>x</p></body></html>", encoding="utf-8")

    doc = parse_document(f)
    assert doc.meta["title"] == "plain"


def test_unknown_extension_falls_back_to_text_read(tmp_path: Path):
    """无内置 parser 声明的扩展名仍走 UTF-8 兜底，且不带 parser 标记。"""
    f = tmp_path / "data.json"
    f.write_text('{"k": "v"}', encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    assert "parser" not in doc.meta
    assert doc.meta["title"] == "data"
    assert doc.format == "json"


def test_missing_and_directory_return_none(tmp_path: Path):
    assert parse_document(tmp_path / "nope.md") is None
    assert parse_document(tmp_path) is None


def test_empty_file_returns_document_without_segments(tmp_path: Path):
    f = tmp_path / "empty.md"
    f.write_text("", encoding="utf-8")

    doc = parse_document(f)
    assert doc is not None
    assert doc.segments == []


# ---------------------------------------------------------------------------
# 兼容 API
# ---------------------------------------------------------------------------


def test_parse_html_compat_api(tmp_path: Path):
    f = tmp_path / "page.html"
    f.write_text("<title>Compat</title><h1>H</h1><p>Body</p>", encoding="utf-8")

    doc = parse_html(f)
    assert doc.format == "html"
    assert doc.meta["title"] == "Compat"
    assert doc.source_path == str(f)


def test_list_supported_formats_merges_builtin_and_custom():
    def custom(path: Path) -> ParsedDocument:
        return ParsedDocument(text="x")

    register_parser([".bridgefmt"], custom)

    formats = list_supported_formats()

    # 内置 registry 声明的扩展名
    for ext in (".md", ".html", ".pdf", ".docx", ".xlsx", ".pptx", ".png", ".py"):
        assert ext in formats
    # 外部注册的扩展名
    assert ".bridgefmt" in formats
    # 无扩展名不是可枚举的“格式”
    assert "" not in formats


@pytest.mark.parametrize(
    "name,expected",
    [
        ("a.mdx", "markdown"),
        ("a.xhtml", "html"),
        ("a.pyi", "python"),
        ("a.cs", "csharp"),
        ("a.kt", "kotlin"),
        ("a.sql", "sql"),
        ("a.pdf", "pdf"),
        ("a.docx", "docx"),
        ("a.xlsx", "xlsx"),
        ("a.xlsm", "xlsx"),
        ("a.pptx", "pptx"),
        ("a.png", "image"),
        ("a.jpeg", "image"),
    ],
)
def test_detect_format_covers_builtin_extensions(name: str, expected: str):
    assert detect_format(Path("/tmp") / name) == expected


# ---------------------------------------------------------------------------
# 二进制格式：旧入口必须拿到 registry 的能力（旧实现只会当纯文本读）
# ---------------------------------------------------------------------------


def test_pdf_uses_pdf_parser(tmp_path: Path):
    pypdf = pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    f = tmp_path / "doc.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(f, "wb") as fh:
        writer.write(fh)

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "pdf"
    assert doc.format == "pdf"
    assert doc.meta["page_count"] == 1
    assert doc.source_path == str(f)


def test_docx_uses_docx_parser(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document as DocxDocument

    f = tmp_path / "doc.docx"
    docx = DocxDocument()
    docx.add_paragraph("Hello DOCX")
    docx.add_paragraph("Second paragraph")
    docx.save(str(f))

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "docx"
    assert doc.format == "docx"
    assert "Hello DOCX" in doc.text
    assert doc.meta["paragraph_count"] == 2


def test_xlsx_uses_xlsx_parser(tmp_path: Path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    f = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["name", "amount"])
    ws.append(["alice", "100"])
    wb.save(str(f))

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "xlsx"
    assert doc.format == "xlsx"
    assert doc.meta["row_count"] == 2
    assert "alice | 100" in doc.text


def test_pptx_uses_pptx_parser(tmp_path: Path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    f = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(8), Inches(5)
    ).text_frame.text = "Body line 1"
    prs.save(str(f))

    doc = parse_document(f)
    assert doc is not None
    assert doc.meta["parser"] == "pptx"
    assert doc.format == "pptx"
    assert doc.meta["slide_count"] == 1
    assert "Body line 1" in doc.text


def test_image_without_ocr_returns_none_not_crash(tmp_path: Path):
    """缺 OCR 依赖时图片应返回 None，而不是抛异常或返回乱码文本。"""
    f = tmp_path / "shot.png"
    f.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 40)

    doc = parse_document(f)
    if doc is None:
        pytest.skip("本环境未安装 OCR 依赖")
    assert doc.meta["parser"] == "ocr"
    assert doc.format == "image"


# ---------------------------------------------------------------------------
# 文件发现：CLI 与 MCP 必须用同一份扩展名集合
# ---------------------------------------------------------------------------


def test_importable_extensions_cover_registry_and_text_fallback():
    exts = importable_extensions()

    # registry 声明的全部格式（含此前硬编码集合漏掉的）
    for ext in (
        ".md", ".mdx", ".html", ".htm", ".xhtml", ".txt", ".log",
        ".py", ".js", ".ts", ".go", ".rs", ".java", ".cs",
        ".pdf", ".docx", ".xlsx", ".pptx",
        ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff",
    ):
        assert ext in exts, ext

    # 纯文本兜底格式
    for ext in (".json", ".yaml", ".yml", ".xml", ".csv"):
        assert ext in exts, ext

    # 无扩展名不是可枚举的发现目标
    assert "" not in exts


def test_cli_and_mcp_share_extension_source():
    """防止 CLI / MCP 再次各自硬编码。"""
    import inspect

    from knowbase import __main__ as cli
    from knowbase import server

    cli_src = inspect.getsource(cli.cmd_import)
    mcp_src = inspect.getsource(server.import_impl)

    assert "importable_extensions()" in cli_src
    assert "importable_extensions()" in mcp_src
    assert ".pdf" not in cli_src  # 不允许再出现硬编码扩展名集合