"""标注完成后的一键复核：读填写后的 labeling-sheet.xlsx，重算 A/B 并输出 P@3 与 Phase 1 终判。

用法：两位置标注人分别填"待标注"表（labeler 列署名），本脚本可对任一版本运行：
  /Users/qc/work/knowbase-phase1/.venv-p1/bin/python eval/compute_p3.py <填好的xlsx路径>
"""

import json
import os
import sys
import time
from pathlib import Path

sys.path.insert(0, "/Users/qc/work/knowbase-phase1")
os.environ["KNOWBASE_REPO_PATH"] = "/tmp/kb-phase1-eval"

from openpyxl import load_workbook  # noqa: E402
from knowbase import index, query_hygiene  # noqa: E402

XLSX = Path(sys.argv[1] if len(sys.argv) > 1 else
           "/Users/qc/work/knowbase/docs/retrieval-upgrade/eval/labeling-sheet.xlsx")

wb = load_workbook(XLSX, data_only=True)
ws = wb["待标注"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
labeled = []
for r in rows:
    if not r or not r[0]:
        continue
    rid, stratum, _, _, query, raw_len, _, _, rel, unans, confusable, labeler, _notes = r[:13]
    rel_ids = [t.strip() for t in str(rel or "").replace("；", ";").split(";") if t.strip()]
    if rel_ids or str(unans).strip() in ("1", "是", "y", "Y"):
        labeled.append({"id": str(rid).strip(), "stratum": stratum, "query": str(query),
                        "raw_len": int(raw_len or len(str(query))),
                        "relevant": rel_ids,
                        "unanswerable": str(unans).strip() in ("1", "是", "y", "Y"),
                        "labeler": labeler})

if not labeled:
    sys.exit("未读到任何标注行（relevant_ids 或 unanswerable 至少填一项）")
print(f"读到标注 {len(labeled)} 条；labelers={sorted({r['labeler'] for r in labeled})}")

cands = {json.loads(l)["id"]: json.loads(l) for l in
         Path("/Users/qc/work/knowbase/docs/retrieval-upgrade/eval/candidates.jsonl").read_text().splitlines() if l.strip()}


def hook_like(conn, query, scope):
    hits = index.search(conn, query, scope=scope, limit=100)
    return [h for h in hits if h.get("confidence") == "verified"][:3]


conn = index.connect(Path("/tmp/kb-phase1-eval"))
stats = {"A": {"p3": [], "rej_ok": [], "rej_n": 0}, "B": {"p3": [], "rej_ok": [], "rej_n": 0}}
detail = []
for r in labeled:
    c = cands.get(r["id"], {})
    scope = c.get("scope")
    a = [h["id"] for h in hook_like(conn, r["query"], scope)]
    probe = (query_hygiene.prepare(r["query"]) if r["raw_len"] <= 512
             else {"action": "skip", "reason": "too_long"})
    b = [] if probe["action"] == "skip" else [h["id"] for h in hook_like(conn, probe["query"], scope)]
    if r["unanswerable"]:
        for arm, hits in (("A", a), ("B", b)):
            stats[arm]["rej_n"] += 1
            stats[arm]["rej_ok"].append(1 if not hits else 0)
    elif r["relevant"]:
        for arm, hits in (("A", a), ("B", b)):
            hit = sum(1 for h in hits if h in r["relevant"])
            stats[arm]["p3"].append(hit / 3.0)
        if a and not b:
            detail.append(("B零注入(待仲裁)", r["id"], r["query"][:40], a, r["relevant"]))
        elif b and not a:
            detail.append(("B新增注入", r["id"], r["query"][:40], b, r["relevant"]))
conn.close()

out = {"labeled": len(labeled), "labelers": sorted({r["labeler"] for r in labeled})}
for arm in ("A", "B"):
    p3 = stats[arm]["p3"]
    rej = stats[arm]["rej_ok"]
    out[arm] = {"P@3": round(sum(p3) / len(p3), 3) if p3 else None,
                "P@3_n": len(p3),
                "正确拒绝率": f"{sum(rej)}/{len(rej)}" if rej else "n/a"}
delta = None
if out["A"]["P@3"] is not None and out["B"]["P@3"] is not None:
    delta = round(out["B"]["P@3"] - out["A"]["P@3"], 3)
out["P@3_delta(B-A)"] = delta
out["待仲裁"] = detail
print(json.dumps(out, ensure_ascii=False, indent=1))
p3_ok = delta is not None and delta >= 0
rej_a = sum(stats["A"]["rej_ok"]) if stats["A"]["rej_ok"] else None
rej_b = sum(stats["B"]["rej_ok"]) if stats["B"]["rej_ok"] else None
rej_ok = rej_a is None or (rej_b is not None and rej_b >= rej_a)
verdict = p3_ok and rej_ok
print("\nPhase 1 P@3 门槛判定：", "通过（B≥A 且正确拒绝不退）" if verdict else "未通过——见待仲裁清单，仲裁后重跑")
