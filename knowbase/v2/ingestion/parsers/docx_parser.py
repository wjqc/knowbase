"""DOCX 解析器（P1-B）。

基于 python-docx：
- 按段落顺序提取正文；段间用 \n
- meta 记录段落数与 title（首段 / 文件 stem）
- 表格内容用「列分隔 | 行分隔 \n」输出，便于检索
"""
from __future__ import annotations

from pathlib import Path

from .base import DocumentParser, ParsedDocument


class DocxParser(DocumentParser):
    name = "docx"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def parse(self, path: Path) -> ParsedDocument:
        from ...observability.errors import ParseError
        try:
            from docx import Document as DocxDocument  # type: ignore
        except ImportError as e:
            raise ParseError(parser=self.name,
                             reason="python-docx not installed; uv add python-docx",
                             path=str(path)) from e
        try:
            doc = DocxDocument(str(path))
        except Exception as e:
            raise ParseError(parser=self.name, reason=f"cannot open: {e}",
                             path=str(path)) from e

        parts: list[str] = []
        for p in doc.paragraphs:
            if p.text:
                parts.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))

        text = "\n".join(parts).strip()
        title = doc.paragraphs[0].text.strip()[:80] if doc.paragraphs and doc.paragraphs[0].text else path.stem
        return ParsedDocument(
            text=text,
            meta={"format": "docx", "title": title, "paragraph_count": len(doc.paragraphs)},
        )
