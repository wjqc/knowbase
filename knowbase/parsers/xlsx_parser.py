"""XLSX 解析器（P5-A1）。

基于 openpyxl：
- 每个 sheet：表头行（首行）+ 数据行；行用 ` | ` 分隔列，sheet 间用 `\n\n`
- meta 记录 sheet_count / sheet_names / row_count / title
- 缺失依赖抛 ParseError 提示 `uv add openpyxl`
- 加密 / 损坏文件抛 ParseError（reason 含原因）
- 空 sheet / 全空行跳过
"""
from __future__ import annotations

from pathlib import Path

from .base import DocumentParser, ParsedDocument


class XlsxParser(DocumentParser):
    name = "xlsx"

    def supports(self, path: Path) -> bool:
        s = path.suffix.lower()
        return s == ".xlsx" or s == ".xlsm"

    def parse(self, path: Path) -> ParsedDocument:
        from .errors import ParseError
        try:
            import openpyxl  # type: ignore
        except ImportError as e:
            raise ParseError(parser=self.name,
                             reason="openpyxl not installed; uv add openpyxl",
                             path=str(path)) from e
        try:
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        except Exception as e:
            raise ParseError(parser=self.name, reason=f"cannot open: {e}",
                             path=str(path)) from e

        sheets_text: list[str] = []
        sheet_names: list[str] = []
        total_rows = 0
        try:
            for ws in wb.worksheets:
                sheet_names.append(ws.title)
                lines: list[str] = []
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v).strip() for v in row]
                    if any(cells):
                        lines.append(" | ".join(cells))
                        total_rows += 1
                if lines:
                    sheets_text.append(f"# Sheet: {ws.title}\n" + "\n".join(lines))
        finally:
            try:
                wb.close()
            except Exception:
                pass

        text = "\n\n".join(sheets_text).strip()
        title = sheet_names[0] if sheet_names else path.stem
        return ParsedDocument(
            text=text,
            meta={
                "format": "xlsx",
                "title": title,
                "sheet_count": len(sheet_names),
                "sheet_names": sheet_names,
                "row_count": total_rows,
            },
        )